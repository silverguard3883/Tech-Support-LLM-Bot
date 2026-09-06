from dataclasses import dataclass
from pathlib import Path

from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader

try:
    import trafilatura
except Exception:  # pragma: no cover - optional extraction helper
    trafilatura = None


@dataclass
class ParsedDocument:
    """Document text and basic metadata extracted from a file."""

    title: str
    text: str
    metadata: dict


def parse_file(path: Path) -> ParsedDocument:
    """Parse supported document formats into plain text."""
    suffix = path.suffix.lower()

    if suffix in {".txt", ".md", ".csv"}:
        return _parse_plain_text(path)
    if suffix == ".docx":
        return _parse_docx(path)
    if suffix == ".pdf":
        return _parse_pdf(path)
    if suffix in {".html", ".htm"}:
        return _parse_html(path)

    raise ValueError(f"Unsupported file type: {path.suffix}")


def _parse_plain_text(path: Path) -> ParsedDocument:
    """Read text-like files directly."""
    text = path.read_text(encoding="utf-8", errors="ignore")
    return ParsedDocument(path.stem, text, {"parser": "plain_text"})


def _parse_docx(path: Path) -> ParsedDocument:
    """Extract paragraphs and table cells from Word documents."""
    doc = Document(str(path))
    pieces = []

    for paragraph in doc.paragraphs:
        if paragraph.text.strip():
            pieces.append(paragraph.text.strip())

    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                pieces.append(" | ".join(cells))

    title = doc.core_properties.title or path.stem
    return ParsedDocument(title, "\n\n".join(pieces), {"parser": "python-docx"})


def _parse_pdf(path: Path) -> ParsedDocument:
    """Extract text from PDF pages without OCR."""
    reader = PdfReader(str(path))
    pages = []

    for page_number, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"[Page {page_number}]\n{text.strip()}")

    metadata = {
        "parser": "pypdf",
        "pages": len(reader.pages),
    }
    return ParsedDocument(path.stem, "\n\n".join(pages), metadata)


def _parse_html(path: Path) -> ParsedDocument:
    """Extract readable text from HTML files."""
    html = path.read_text(encoding="utf-8", errors="ignore")
    return parse_html_text(html, fallback_title=path.stem)


def parse_html_text(html: str, fallback_title: str = "Untitled Page") -> ParsedDocument:
    """Extract main text from HTML using Trafilatura first, then BeautifulSoup."""
    title = fallback_title
    extracted = ""

    if trafilatura is not None:
        extracted = trafilatura.extract(html) or ""

    soup = BeautifulSoup(html, "html.parser")
    if soup.title and soup.title.string:
        title = soup.title.string.strip()

    if not extracted:
        for item in soup(["script", "style", "noscript"]):
            item.decompose()
        extracted = soup.get_text("\n", strip=True)

    return ParsedDocument(title, extracted, {"parser": "html"})
